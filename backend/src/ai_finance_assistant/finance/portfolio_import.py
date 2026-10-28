"""One safe import interface over manual, JSON, CSV, and XLSX representations."""

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import defaultdict
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from ai_finance_assistant.contracts.portfolio import SecurityIdentifier, SecurityIdentifierType
from ai_finance_assistant.finance.models import (
    CsvImportPayload,
    FindingScope,
    FindingSeverity,
    ImportedAccount,
    ImportedHolding,
    ImportedPortfolio,
    ImportFinding,
    JsonImportPayload,
    ManualImportPayload,
    PortfolioImportRequest,
    PortfolioImportResult,
    RawFieldValue,
    XlsxImportPayload,
)

IMPORT_POLICY_VERSION = "portfolio-import-v1"
MAX_FILE_BYTES = 2_000_000
MAX_DECOMPRESSED_BYTES = 20_000_000
MAX_ROWS = 10_000
MAX_COLUMNS = 32
MAX_TEXT_LENGTH = 512
_UNSAFE_REPLACEMENT = "[unsafe cell content removed]"
_NUMERIC_TEXT = re.compile(r"^-?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
_IDENTIFIER_PATTERNS = {
    SecurityIdentifierType.TICKER: re.compile(r"^[A-Z0-9][A-Z0-9.\-]{0,31}$"),
    SecurityIdentifierType.CUSIP: re.compile(r"^[A-Z0-9]{9}$"),
    SecurityIdentifierType.ISIN: re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$"),
    SecurityIdentifierType.USER_DEFINED: re.compile(r"^.{1,64}$"),
}
_REQUIRED_FIELDS = {
    "account_id",
    "account_type",
    "holding_id",
    "identifier_type",
    "identifier_value",
    "quantity",
    "currency",
}
_SUPPORTED_FIELDS = _REQUIRED_FIELDS | {
    "institution_label",
    "asset_class",
    "cost_basis_amount",
    "cost_basis_basis",
    "acquisition_date",
    "user_price",
    "user_price_as_of",
    "user_price_provenance",
    "expense_ratio",
    "sector",
}
_HEADER_ALIASES = {
    "account": "account_id",
    "account id": "account_id",
    "account type": "account_type",
    "institution": "institution_label",
    "institution label": "institution_label",
    "holding": "holding_id",
    "holding id": "holding_id",
    "identifier type": "identifier_type",
    "identifier": "identifier_value",
    "identifier value": "identifier_value",
    "asset class": "asset_class",
    "cost basis": "cost_basis_amount",
    "cost basis amount": "cost_basis_amount",
    "cost basis basis": "cost_basis_basis",
    "acquisition date": "acquisition_date",
    "user price": "user_price",
    "user price as of": "user_price_as_of",
    "user price provenance": "user_price_provenance",
    "expense ratio": "expense_ratio",
}


@dataclass(frozen=True)
class _SourceRow:
    number: int
    values: Mapping[str, object]


def _finding(
    code: str,
    severity: FindingSeverity,
    scope: FindingScope,
    message: str,
    *,
    row: int | None = None,
    field: str | None = None,
) -> ImportFinding:
    return ImportFinding(
        code=code,
        severity=severity,
        scope=scope,
        message=message,
        row_number=row,
        field=field,
    )


def _normal_header(value: object) -> str:
    text = str(value).strip().casefold().replace("_", " ").replace("-", " ")
    text = " ".join(text.split())
    return _HEADER_ALIASES.get(text, text.replace(" ", "_"))


def _looks_unsafe(value: object) -> bool:
    if not isinstance(value, str):
        return False
    stripped = value.lstrip()
    if not stripped:
        return False
    if stripped.startswith("-") and _NUMERIC_TEXT.fullmatch(stripped):
        return False
    return stripped[0] in {"=", "+", "-", "@"}


def _safe_raw(value: object) -> str | None:
    if value is None:
        return None
    if _looks_unsafe(value):
        return _UNSAFE_REPLACEMENT
    text = str(value)
    return text[:MAX_TEXT_LENGTH]


def _source_hash(payload: object) -> str:
    if isinstance(payload, bytes):
        material = payload
    else:
        material = json.dumps(
            payload,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode()
    return hashlib.sha256(material).hexdigest()


def _decode_utf8(content: bytes) -> str:
    return content.decode("utf-8-sig", errors="strict")


def _check_file_basics(content: bytes, filename: str, suffix: str) -> list[ImportFinding]:
    findings: list[ImportFinding] = []
    if Path(filename).suffix.casefold() != suffix:
        findings.append(
            _finding(
                "UNSUPPORTED_FILE_TYPE",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                f"Expected a {suffix} file.",
            )
        )
    if not content:
        findings.append(
            _finding(
                "EMPTY_FILE",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "The import file is empty.",
            )
        )
    if len(content) > MAX_FILE_BYTES:
        findings.append(
            _finding(
                "FILE_SIZE_LIMIT",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                f"The import exceeds the {MAX_FILE_BYTES}-byte limit.",
            )
        )
    return findings


def _manual_rows(payload: ManualImportPayload) -> tuple[list[_SourceRow], list[ImportFinding]]:
    if len(payload.rows) > MAX_ROWS:
        return [], [
            _finding(
                "ROW_LIMIT",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                f"The import exceeds the {MAX_ROWS}-row limit.",
            )
        ]
    if any(len(row) > MAX_COLUMNS for row in payload.rows):
        return [], [
            _finding(
                "COLUMN_LIMIT",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                f"The import exceeds the {MAX_COLUMNS}-column limit.",
            )
        ]
    return [
        _SourceRow(number=index, values=row) for index, row in enumerate(payload.rows, start=1)
    ], []


def _json_rows(payload: JsonImportPayload) -> tuple[list[_SourceRow], list[ImportFinding]]:
    findings = _check_file_basics(payload.content, payload.filename, ".json")
    if findings:
        return [], findings
    try:
        raw = json.loads(
            _decode_utf8(payload.content),
            parse_float=Decimal,
            parse_int=Decimal,
        )
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        return [], [
            _finding(
                "INVALID_JSON",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                (
                    "The JSON document is invalid at a reported parser position: "
                    f"{error.pos if isinstance(error, json.JSONDecodeError) else 'encoding'}."
                ),
            )
        ]
    if isinstance(raw, dict) and set(raw) == {"rows"}:
        raw = raw["rows"]
    if not isinstance(raw, list):
        return [], [
            _finding(
                "INVALID_JSON_SHAPE",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "JSON must contain a list of holding rows or an object with only a rows list.",
            )
        ]
    if len(raw) > MAX_ROWS:
        return [], [
            _finding(
                "ROW_LIMIT",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                f"The import exceeds the {MAX_ROWS}-row limit.",
            )
        ]
    rows: list[_SourceRow] = []
    for index, item in enumerate(raw, start=1):
        if not isinstance(item, dict) or not all(isinstance(key, str) for key in item):
            findings.append(
                _finding(
                    "INVALID_ROW_SHAPE",
                    FindingSeverity.ERROR,
                    FindingScope.ROW,
                    "Each JSON row must be an object with string field names.",
                    row=index,
                )
            )
            continue
        if len(item) > MAX_COLUMNS:
            findings.append(
                _finding(
                    "COLUMN_LIMIT",
                    FindingSeverity.FATAL,
                    FindingScope.FILE,
                    f"The import exceeds the {MAX_COLUMNS}-column limit.",
                )
            )
            continue
        rows.append(_SourceRow(number=index, values=item))
    return rows, findings


def _csv_rows(payload: CsvImportPayload) -> tuple[list[_SourceRow], list[ImportFinding]]:
    findings = _check_file_basics(payload.content, payload.filename, ".csv")
    if findings:
        return [], findings
    try:
        text = _decode_utf8(payload.content)
    except UnicodeDecodeError:
        return [], [
            _finding(
                "INVALID_ENCODING",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "CSV must use valid UTF-8 encoding.",
            )
        ]
    if "\x00" in text:
        return [], [
            _finding(
                "NUL_BYTE",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "CSV contains a prohibited NUL byte.",
            )
        ]
    try:
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        rows = list(reader)
    except csv.Error:
        return [], [
            _finding(
                "INVALID_CSV",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "CSV structure is invalid.",
            )
        ]
    if not rows:
        return [], [
            _finding(
                "EMPTY_FILE",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "The import file contains no rows.",
            )
        ]
    if len(rows) - 1 > MAX_ROWS:
        return [], [
            _finding(
                "ROW_LIMIT",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                f"The import exceeds the {MAX_ROWS}-row limit.",
            )
        ]
    if any(len(row) > MAX_COLUMNS for row in rows):
        return [], [
            _finding(
                "COLUMN_LIMIT",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                f"The import exceeds the {MAX_COLUMNS}-column limit.",
            )
        ]
    headers = [_normal_header(value) for value in rows[0]]
    if len(headers) != len(set(headers)):
        findings.append(
            _finding(
                "DUPLICATE_HEADER",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "Headers must be unique after normalization.",
            )
        )
        return [], findings
    parsed: list[_SourceRow] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if len(values) > len(headers):
            findings.append(
                _finding(
                    "EXTRA_COLUMNS",
                    FindingSeverity.ERROR,
                    FindingScope.ROW,
                    "The row has more values than the header.",
                    row=row_number,
                )
            )
            continue
        padded = [*values, *([""] * (len(headers) - len(values)))]
        parsed.append(_SourceRow(row_number, dict(zip(headers, padded, strict=True))))
    return parsed, findings


def _xlsx_rows(payload: XlsxImportPayload) -> tuple[list[_SourceRow], list[ImportFinding]]:
    findings = _check_file_basics(payload.content, payload.filename, ".xlsx")
    if findings:
        return [], findings
    try:
        with zipfile.ZipFile(io.BytesIO(payload.content)) as archive:
            entries = archive.infolist()
            if len(entries) > 200:
                raise ValueError("too many archive members")
            total_uncompressed = sum(entry.file_size for entry in entries)
            if total_uncompressed > MAX_DECOMPRESSED_BYTES:
                raise ValueError("decompressed size limit exceeded")
            lowered = {entry.filename.casefold() for entry in entries}
            if any(
                "vbaproject" in name
                or name.startswith("xl/externallinks/")
                or name.startswith("xl/embeddings/")
                or name.startswith("customui/")
                for name in lowered
            ):
                raise ValueError("macro, external-link, or embedded content is not allowed")
            if any(
                entry.filename.startswith("/")
                or "\\" in entry.filename
                or ".." in Path(entry.filename).parts
                for entry in entries
            ):
                raise ValueError("unsafe archive member path")
            if any(
                entry.filename.casefold().endswith(".rels")
                and b'TargetMode="External"' in archive.read(entry)
                for entry in entries
            ):
                raise ValueError("external workbook relationship is not allowed")
            corrupt = archive.testzip()
            if corrupt is not None:
                raise ValueError("archive integrity check failed")
    except (zipfile.BadZipFile, ValueError):
        return [], [
            _finding(
                "UNSAFE_XLSX_CONTAINER",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                (
                    "XLSX must be a bounded workbook without macros, external links, or "
                    "embedded objects."
                ),
            )
        ]

    try:
        workbook = load_workbook(
            io.BytesIO(payload.content), read_only=True, data_only=False, keep_links=False
        )
    except Exception:
        return [], [
            _finding(
                "INVALID_XLSX",
                FindingSeverity.FATAL,
                FindingScope.FILE,
                "The XLSX workbook could not be parsed safely.",
            )
        ]
    try:
        if len(workbook.worksheets) != 1:
            return [], [
                _finding(
                    "XLSX_SHEET_COUNT",
                    FindingSeverity.FATAL,
                    FindingScope.FILE,
                    "XLSX imports must contain exactly one worksheet.",
                )
            ]
        sheet = workbook.worksheets[0]
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
            return [], [
                _finding(
                    "XLSX_DIMENSION_LIMIT",
                    FindingSeverity.FATAL,
                    FindingScope.FILE,
                    "The XLSX worksheet exceeds the configured row or column limit.",
                )
            ]
        iterator = sheet.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration:
            return [], [
                _finding(
                    "EMPTY_FILE",
                    FindingSeverity.FATAL,
                    FindingScope.FILE,
                    "The XLSX worksheet is empty.",
                )
            ]
        headers = [_normal_header(cell.value or "") for cell in header_cells]
        if not all(headers) or len(headers) != len(set(headers)):
            return [], [
                _finding(
                    "INVALID_HEADER",
                    FindingSeverity.FATAL,
                    FindingScope.FILE,
                    "XLSX headers must be nonempty and unique after normalization.",
                )
            ]
        parsed: list[_SourceRow] = []
        for row_number, cells in enumerate(iterator, start=2):
            if row_number - 1 > MAX_ROWS or len(cells) > MAX_COLUMNS:
                return [], [
                    _finding(
                        "XLSX_DIMENSION_LIMIT",
                        FindingSeverity.FATAL,
                        FindingScope.FILE,
                        "The XLSX worksheet exceeds the configured row or column limit.",
                    )
                ]
            values: dict[str, object] = {}
            row_has_value = False
            for header, cell in zip(headers, cells, strict=False):
                if cell.data_type == "f":
                    findings.append(
                        _finding(
                            "UNSAFE_FORMULA",
                            FindingSeverity.ERROR,
                            FindingScope.FIELD,
                            "Spreadsheet formulas are not accepted as portfolio data.",
                            row=row_number,
                            field=header,
                        )
                    )
                    values[header] = _UNSAFE_REPLACEMENT
                    row_has_value = True
                    continue
                values[header] = cell.value
                row_has_value = row_has_value or cell.value not in {None, ""}
            if row_has_value:
                parsed.append(_SourceRow(row_number, values))
        return parsed, findings
    finally:
        workbook.close()


def _parse_required_text(
    value: object,
    *,
    row: int,
    field: str,
    findings: list[ImportFinding],
    max_length: int = 128,
) -> str | None:
    if value is None or not str(value).strip():
        findings.append(
            _finding(
                "MISSING_REQUIRED_FIELD",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "A required field is missing.",
                row=row,
                field=field,
            )
        )
        return None
    if _looks_unsafe(value):
        findings.append(
            _finding(
                "UNSAFE_FORMULA",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Formula-like cell content is not accepted.",
                row=row,
                field=field,
            )
        )
        return None
    result = str(value).strip()
    if len(result) > max_length:
        findings.append(
            _finding(
                "TEXT_LENGTH_LIMIT",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                f"Text exceeds the {max_length}-character field limit.",
                row=row,
                field=field,
            )
        )
        return None
    return result


def _optional_text(
    value: object,
    *,
    row: int,
    field: str,
    findings: list[ImportFinding],
    max_length: int = 128,
) -> str | None:
    if value is None or not str(value).strip():
        return None
    return _parse_required_text(
        value, row=row, field=field, findings=findings, max_length=max_length
    )


def _parse_decimal(
    value: object,
    *,
    row: int,
    field: str,
    findings: list[ImportFinding],
    required: bool,
) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            findings.append(
                _finding(
                    "MISSING_REQUIRED_FIELD",
                    FindingSeverity.ERROR,
                    FindingScope.FIELD,
                    "A required numeric field is missing.",
                    row=row,
                    field=field,
                )
            )
        return None
    if _looks_unsafe(value):
        findings.append(
            _finding(
                "UNSAFE_FORMULA",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Formula-like numeric content is not accepted.",
                row=row,
                field=field,
            )
        )
        return None
    if isinstance(value, (float, bool)):
        findings.append(
            _finding(
                "INEXACT_NUMERIC_TYPE",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Use a decimal string or integer for exact numeric import.",
                row=row,
                field=field,
            )
        )
        return None
    try:
        parsed = value if isinstance(value, Decimal) else Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        findings.append(
            _finding(
                "INVALID_DECIMAL",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "The field must contain a finite decimal number.",
                row=row,
                field=field,
            )
        )
        return None
    if not parsed.is_finite():
        findings.append(
            _finding(
                "NONFINITE_NUMBER",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Nonfinite numbers are prohibited.",
                row=row,
                field=field,
            )
        )
        return None
    return parsed


def _parse_date(
    value: object,
    *,
    row: int,
    field: str,
    findings: list[ImportFinding],
) -> date | None:
    text = _optional_text(value, row=row, field=field, findings=findings)
    if text is None:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        findings.append(
            _finding(
                "INVALID_DATE",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Dates must use ISO YYYY-MM-DD format.",
                row=row,
                field=field,
            )
        )
        return None


def _parse_datetime(
    value: object,
    *,
    row: int,
    field: str,
    findings: list[ImportFinding],
) -> datetime | None:
    text = _optional_text(value, row=row, field=field, findings=findings)
    if text is None:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        parsed = None
    if parsed is None or parsed.tzinfo is None or parsed.utcoffset() is None:
        findings.append(
            _finding(
                "INVALID_TIMESTAMP",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Timestamps must be timezone-aware ISO-8601 values.",
                row=row,
                field=field,
            )
        )
        return None
    return parsed


def _build_holding(
    row: _SourceRow, findings: list[ImportFinding]
) -> tuple[ImportedHolding | None, str | None, str | None]:
    finding_start = len(findings)
    values = {_normal_header(key): value for key, value in row.values.items()}
    for name in sorted(set(values).difference(_SUPPORTED_FIELDS)):
        findings.append(
            _finding(
                "UNKNOWN_FIELD",
                FindingSeverity.WARNING,
                FindingScope.FIELD,
                "An unrecognized field was ignored.",
                row=row.number,
                field=name[:64],
            )
        )
    account_id = _parse_required_text(
        values.get("account_id"), row=row.number, field="account_id", findings=findings
    )
    account_type = _parse_required_text(
        values.get("account_type"), row=row.number, field="account_type", findings=findings
    )
    holding_id = _parse_required_text(
        values.get("holding_id"), row=row.number, field="holding_id", findings=findings
    )
    identifier_type_text = _parse_required_text(
        values.get("identifier_type"),
        row=row.number,
        field="identifier_type",
        findings=findings,
    )
    identifier_value = _parse_required_text(
        values.get("identifier_value"),
        row=row.number,
        field="identifier_value",
        findings=findings,
        max_length=64,
    )
    quantity = _parse_decimal(
        values.get("quantity"),
        row=row.number,
        field="quantity",
        findings=findings,
        required=True,
    )
    currency = _parse_required_text(
        values.get("currency"),
        row=row.number,
        field="currency",
        findings=findings,
        max_length=3,
    )

    identifier_type: SecurityIdentifierType | None = None
    if identifier_type_text is not None:
        try:
            identifier_type = SecurityIdentifierType(identifier_type_text.casefold())
        except ValueError:
            findings.append(
                _finding(
                    "UNSUPPORTED_IDENTIFIER_TYPE",
                    FindingSeverity.ERROR,
                    FindingScope.FIELD,
                    "Identifier type must be ticker, cusip, isin, or user_defined.",
                    row=row.number,
                    field="identifier_type",
                )
            )
    if identifier_type is not None and identifier_value is not None:
        identifier_value = identifier_value.strip()
        if identifier_type != SecurityIdentifierType.USER_DEFINED:
            identifier_value = identifier_value.upper()
        if not _IDENTIFIER_PATTERNS[identifier_type].fullmatch(identifier_value):
            findings.append(
                _finding(
                    "INVALID_IDENTIFIER",
                    FindingSeverity.ERROR,
                    FindingScope.FIELD,
                    "The identifier does not match its declared type.",
                    row=row.number,
                    field="identifier_value",
                )
            )

    if quantity is not None and quantity < 0:
        findings.append(
            _finding(
                "SHORT_POSITION_UNSUPPORTED",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Negative quantities are unsupported in calculation policy v1.",
                row=row.number,
                field="quantity",
            )
        )
    elif quantity == 0:
        findings.append(
            _finding(
                "ZERO_QUANTITY",
                FindingSeverity.WARNING,
                FindingScope.FIELD,
                "A zero-quantity Holding is preserved and valued at zero when priced.",
                row=row.number,
                field="quantity",
            )
        )
    if currency is not None and not re.fullmatch(r"[A-Za-z]{3}", currency):
        findings.append(
            _finding(
                "INVALID_CURRENCY",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Currency must be a three-letter code.",
                row=row.number,
                field="currency",
            )
        )

    cost_basis_amount = _parse_decimal(
        values.get("cost_basis_amount"),
        row=row.number,
        field="cost_basis_amount",
        findings=findings,
        required=False,
    )
    cost_basis_basis = _optional_text(
        values.get("cost_basis_basis"),
        row=row.number,
        field="cost_basis_basis",
        findings=findings,
    )
    user_price = _parse_decimal(
        values.get("user_price"),
        row=row.number,
        field="user_price",
        findings=findings,
        required=False,
    )
    expense_ratio = _parse_decimal(
        values.get("expense_ratio"),
        row=row.number,
        field="expense_ratio",
        findings=findings,
        required=False,
    )
    for field_name, decimal_value, allow_zero in (
        ("cost_basis_amount", cost_basis_amount, True),
        ("user_price", user_price, False),
        ("expense_ratio", expense_ratio, True),
    ):
        if decimal_value is not None and (
            decimal_value < 0 or (not allow_zero and decimal_value == 0)
        ):
            findings.append(
                _finding(
                    "OUT_OF_RANGE_NUMBER",
                    FindingSeverity.ERROR,
                    FindingScope.FIELD,
                    "The numeric value is outside its allowed range.",
                    row=row.number,
                    field=field_name,
                )
            )
    if expense_ratio is not None and expense_ratio > 1:
        findings.append(
            _finding(
                "OUT_OF_RANGE_NUMBER",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Expense ratio must be expressed as a decimal fraction from 0 through 1.",
                row=row.number,
                field="expense_ratio",
            )
        )
    if (cost_basis_amount is None) != (cost_basis_basis is None):
        findings.append(
            _finding(
                "INCOMPLETE_COST_BASIS",
                FindingSeverity.ERROR,
                FindingScope.ROW,
                "Cost-basis amount and basis must be supplied together.",
                row=row.number,
            )
        )
    if cost_basis_basis not in {None, "total", "per_unit"}:
        findings.append(
            _finding(
                "INVALID_COST_BASIS",
                FindingSeverity.ERROR,
                FindingScope.FIELD,
                "Cost-basis basis must be total or per_unit.",
                row=row.number,
                field="cost_basis_basis",
            )
        )

    user_price_as_of = _parse_datetime(
        values.get("user_price_as_of"),
        row=row.number,
        field="user_price_as_of",
        findings=findings,
    )
    user_price_provenance = _optional_text(
        values.get("user_price_provenance"),
        row=row.number,
        field="user_price_provenance",
        findings=findings,
    )
    if user_price is None and (user_price_as_of is not None or user_price_provenance is not None):
        findings.append(
            _finding(
                "ORPHAN_PRICE_METADATA",
                FindingSeverity.ERROR,
                FindingScope.ROW,
                "User price metadata requires a user price.",
                row=row.number,
            )
        )

    if user_price is None:
        findings.append(
            _finding(
                "MISSING_PRICE",
                FindingSeverity.WARNING,
                FindingScope.ROW,
                "No user price was supplied; valuation requires a separate quote.",
                row=row.number,
            )
        )
    asset_class = _optional_text(
        values.get("asset_class"),
        row=row.number,
        field="asset_class",
        findings=findings,
    )
    sector = _optional_text(values.get("sector"), row=row.number, field="sector", findings=findings)
    if asset_class is None or sector is None:
        findings.append(
            _finding(
                "MISSING_CLASSIFICATION",
                FindingSeverity.WARNING,
                FindingScope.ROW,
                "One or more optional classifications are missing.",
                row=row.number,
            )
        )

    blocking_for_row = any(
        finding.row_number == row.number
        and finding.severity in {FindingSeverity.ERROR, FindingSeverity.FATAL}
        for finding in findings[finding_start:]
    )
    if (
        blocking_for_row
        or account_id is None
        or account_type is None
        or holding_id is None
        or identifier_type is None
        or identifier_value is None
        or quantity is None
        or currency is None
    ):
        return (
            None,
            account_type,
            _optional_text(
                values.get("institution_label"),
                row=row.number,
                field="institution_label",
                findings=findings,
            ),
        )

    raw_fields = tuple(
        RawFieldValue(name=name[:64], value=_safe_raw(value))
        for name, value in sorted(values.items())
        if name in _SUPPORTED_FIELDS
    )
    try:
        holding = ImportedHolding(
            holding_id=holding_id,
            account_id=account_id,
            source_row=row.number,
            security=SecurityIdentifier(
                identifier_type=identifier_type,
                value=identifier_value,
            ),
            quantity=quantity,
            currency=currency.upper(),
            asset_class=asset_class,
            cost_basis_amount=cost_basis_amount,
            cost_basis_basis=cost_basis_basis,
            acquisition_date=_parse_date(
                values.get("acquisition_date"),
                row=row.number,
                field="acquisition_date",
                findings=findings,
            ),
            user_price=user_price,
            user_price_as_of=user_price_as_of,
            user_price_provenance=user_price_provenance,
            expense_ratio=expense_ratio,
            sector=sector,
            raw_fields=raw_fields,
        )
    except ValidationError:
        findings.append(
            _finding(
                "INVALID_HOLDING",
                FindingSeverity.ERROR,
                FindingScope.ROW,
                "The normalized Holding violates the versioned domain contract.",
                row=row.number,
            )
        )
        return None, account_type, None
    institution = _optional_text(
        values.get("institution_label"),
        row=row.number,
        field="institution_label",
        findings=findings,
    )
    return holding, account_type, institution


class PortfolioImport:
    """Deep import module: callers see one method and one aggregate result contract."""

    def import_portfolio(self, request: PortfolioImportRequest) -> PortfolioImportResult:
        payload = request.payload
        if isinstance(payload, ManualImportPayload):
            rows, findings = _manual_rows(payload)
            source_hash = _source_hash(payload.rows)
        elif isinstance(payload, JsonImportPayload):
            rows, findings = _json_rows(payload)
            source_hash = _source_hash(payload.content)
        elif isinstance(payload, CsvImportPayload):
            rows, findings = _csv_rows(payload)
            source_hash = _source_hash(payload.content)
        elif isinstance(payload, XlsxImportPayload):
            rows, findings = _xlsx_rows(payload)
            source_hash = _source_hash(payload.content)
        else:  # pragma: no cover - discriminated contract makes this unreachable.
            raise TypeError(f"unsupported import payload {type(payload)!r}")

        if any(finding.severity == FindingSeverity.FATAL for finding in findings):
            return PortfolioImportResult(
                source_kind=payload.kind,
                source_file_hash=source_hash,
                accepted=False,
                portfolio=None,
                findings=tuple(findings),
            )
        if not rows:
            findings.append(
                _finding(
                    "NO_DATA_ROWS",
                    FindingSeverity.FATAL,
                    FindingScope.FILE,
                    "The import contains no usable holding rows.",
                )
            )
            return PortfolioImportResult(
                source_kind=payload.kind,
                source_file_hash=source_hash,
                accepted=False,
                portfolio=None,
                findings=tuple(findings),
            )

        normalized_headers = {_normal_header(key) for row in rows for key in row.values}
        for missing in sorted(_REQUIRED_FIELDS.difference(normalized_headers)):
            findings.append(
                _finding(
                    "MISSING_REQUIRED_HEADER",
                    FindingSeverity.FATAL,
                    FindingScope.FILE,
                    "A required import field is absent from the dataset.",
                    field=missing,
                )
            )
        if any(finding.severity == FindingSeverity.FATAL for finding in findings):
            return PortfolioImportResult(
                source_kind=payload.kind,
                source_file_hash=source_hash,
                accepted=False,
                portfolio=None,
                findings=tuple(findings),
            )

        holdings_by_account: defaultdict[str, list[ImportedHolding]] = defaultdict(list)
        account_metadata: dict[str, tuple[str, str | None]] = {}
        seen_raw_rows: dict[str, int] = {}
        seen_holding_ids: dict[str, int] = {}
        seen_position_keys: dict[tuple[str, SecurityIdentifierType, str], int] = {}

        for row in rows:
            canonical_row = json.dumps(
                {_normal_header(key): str(value) for key, value in row.values.items()},
                sort_keys=True,
                separators=(",", ":"),
            )
            previous_row = seen_raw_rows.get(canonical_row)
            if previous_row is not None:
                findings.append(
                    _finding(
                        "DUPLICATE_ROW",
                        FindingSeverity.ERROR,
                        FindingScope.ROW,
                        f"This row duplicates source row {previous_row}; no merge was performed.",
                        row=row.number,
                    )
                )
            else:
                seen_raw_rows[canonical_row] = row.number

            holding, account_type, institution = _build_holding(row, findings)
            if holding is None:
                continue
            previous_holding = seen_holding_ids.get(holding.holding_id)
            if previous_holding is not None:
                findings.append(
                    _finding(
                        "DUPLICATE_HOLDING_ID",
                        FindingSeverity.ERROR,
                        FindingScope.FIELD,
                        (
                            f"Holding identifier duplicates source row {previous_holding}; "
                            "no merge was performed."
                        ),
                        row=row.number,
                        field="holding_id",
                    )
                )
            else:
                seen_holding_ids[holding.holding_id] = row.number
            position_key = (
                holding.account_id,
                holding.security.identifier_type,
                holding.security.value.casefold(),
            )
            previous_position = seen_position_keys.get(position_key)
            if previous_position is not None:
                findings.append(
                    _finding(
                        "LIKELY_DUPLICATE_HOLDING",
                        FindingSeverity.ERROR,
                        FindingScope.ROW,
                        (
                            "The same account and security appear on source row "
                            f"{previous_position}; an explicit merge decision is required."
                        ),
                        row=row.number,
                    )
                )
            else:
                seen_position_keys[position_key] = row.number

            asserted_metadata = (account_type or "", institution)
            existing_metadata = account_metadata.get(holding.account_id)
            if existing_metadata is not None and existing_metadata != asserted_metadata:
                findings.append(
                    _finding(
                        "ACCOUNT_METADATA_CONFLICT",
                        FindingSeverity.ERROR,
                        FindingScope.ROW,
                        "Rows for the same account contain conflicting account metadata.",
                        row=row.number,
                    )
                )
            else:
                account_metadata[holding.account_id] = asserted_metadata
            holdings_by_account[holding.account_id].append(holding)

        has_blocking = any(
            finding.severity in {FindingSeverity.ERROR, FindingSeverity.FATAL}
            for finding in findings
        )
        if has_blocking:
            return PortfolioImportResult(
                source_kind=payload.kind,
                source_file_hash=source_hash,
                accepted=False,
                portfolio=None,
                findings=tuple(findings),
            )

        accounts = tuple(
            ImportedAccount(
                account_id=account_id,
                portfolio_id=request.metadata.portfolio_id,
                account_type=account_metadata[account_id][0],
                institution_label=account_metadata[account_id][1],
                holdings=tuple(sorted(holdings, key=lambda item: item.holding_id)),
            )
            for account_id, holdings in sorted(holdings_by_account.items())
        )
        portfolio = ImportedPortfolio(
            portfolio_id=request.metadata.portfolio_id,
            label=request.metadata.label,
            base_currency=request.metadata.base_currency,
            imported_at=request.metadata.imported_at,
            source_kind=payload.kind,
            source_file_hash=source_hash,
            accounts=accounts,
        )
        return PortfolioImportResult(
            source_kind=payload.kind,
            source_file_hash=source_hash,
            accepted=True,
            portfolio=portfolio,
            findings=tuple(findings),
        )
